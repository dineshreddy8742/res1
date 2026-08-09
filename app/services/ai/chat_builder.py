"""AIChatResumeBuilder module for interactive conversational resume generation.

This module processes user chat messages, extracts text from uploaded documents
(PDF, DOCX, Images), uses LLMs to ask clarifying questions, and constructs
a complete ATS-optimized ResumeData JSON once sufficient details are gathered.
"""

import json
import logging
import os
import tempfile
from typing import Dict, List, Optional, Tuple, Any

import docx
import PyPDF2
from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage, SystemMessage, AIMessage

from app.database.models.resume import ResumeData

logger = logging.getLogger(__name__)


def extract_file_content(filename: str, content: bytes) -> str:
    """Extract text from uploaded PDF, DOCX, or Image file.

    Args:
        filename: Name of the uploaded file.
        content: Raw bytes of the uploaded file.

    Returns:
        Extracted text string.
    """
    ext = os.path.splitext(filename)[1].lower()
    text = ""

    if ext == ".pdf":
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                tmp.write(content)
                tmp_path = tmp.name

            with open(tmp_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
            os.remove(tmp_path)
        except Exception as e:
            logger.error(f"Error reading PDF {filename}: {e}")
            text = f"[Uploaded PDF document: {filename}]"

    elif ext in [".docx", ".doc"]:
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                tmp.write(content)
                tmp_path = tmp.name

            doc = docx.Document(tmp_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text:
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        full_text.append(row_txt)
            text = "\n".join(full_text)
            os.remove(tmp_path)
        except Exception as e:
            logger.error(f"Error reading DOCX {filename}: {e}")
            text = f"[Uploaded Word document: {filename}]"

    elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
        # Try OCRVision if available, fallback to basic tag
        try:
            from app.utils.vision import OCRVision
            ocr = OCRVision(pdf_bytes=content)
            images = ocr.pdf_to_images()
            if images:
                ocr_texts = []
                for img in images:
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".jpg") as tmp_img:
                        img.save(tmp_img.name, "JPEG")
                        txt = ocr.ocr_image(tmp_img.name)
                        if txt:
                            ocr_texts.append(txt)
                        os.remove(tmp_img.name)
                text = "\n".join(ocr_texts)
        except Exception as e:
            logger.warning(f"OCR not available for image {filename}: {e}")
            text = f"[Uploaded resume image scan: {filename}]"

    if not text.strip():
        text = f"[Uploaded file {filename} contained no readable text]"

    return text.strip()


class AIChatResumeBuilder:
    """Conversational AI Resume Builder service."""

    def __init__(self, model_name: str, api_key: str, api_base: str):
        self.model_name = model_name
        self.api_key = api_key
        self.api_base = api_base
        self.llm = ChatOpenAI(
            model=model_name,
            openai_api_key=api_key,
            openai_api_base=api_base,
            temperature=0.4,
        )

    def process_chat(
        self,
        user_message: str,
        chat_history: List[Dict[str, str]],
        file_info: Optional[Tuple[str, bytes]] = None,
    ) -> Dict[str, Any]:
        """Process chat iteration with optional file attachment.

        Args:
            user_message: Text message from user.
            chat_history: Previous messages list of dicts [{"role": "user"/"assistant", "content": "..."}].
            file_info: Optional tuple of (filename, file_bytes).

        Returns:
            Dict containing:
                - response (str): AI reply message to user
                - is_complete (bool): True if resume generation completed
                - resume_data (Optional[dict]): Parsed ResumeData dict if complete
        """
        extracted_doc_text = ""
        if file_info:
            filename, file_bytes = file_info
            extracted_doc_text = extract_file_content(filename, file_bytes)

        # Build prompt instructions
        system_instruction = (
            "You are AuraRise AI — an elite career strategist and ATS resume architect. "
            "Your goal is to build a top 1% ATS-optimized resume for the user through conversation or by extracting data from their uploaded resume.\n\n"
            "GUIDELINES:\n"
            "1. Be encouraging, concise, and professional.\n"
            "2. Identify what information is already present (Name, Email, Phone, College/Target Role, Experience, Projects, Education, Skills).\n"
            "3. If key essential details are missing and no document was uploaded, ask 1-2 focused questions to gather missing info.\n"
            "4. NEVER say you cannot create PDFs or tell the user to use online JSON converters! AuraRise generates the PDF directly for the user.\n"
            "5. If the user indicates they are a fresher, student, have no work experience/projects, OR when the user asks you to generate/finalize/build the resume, IMMEDIATELY generate the complete JSON resume block!\n\n"
            "CRITICAL PROTOCOL FOR RESUME GENERATION:\n"
            "When ready to finalize, your output MUST end with a single JSON block wrapped in ```json ... ``` containing exact structure:\n"
            "{\n"
            '  "user_information": {\n'
            '    "name": "Full Name",\n'
            '    "email": "user@example.com",\n'
            '    "phone": "+1234567890",\n'
            '    "main_job_title": "Target Job Title",\n'
            '    "profile_description": "High-impact professional summary",\n'
            '    "experiences": [\n'
            '      {"job_title": "Job Title", "company": "Company Name", "location": "City, State", "start_date": "01/2023", "end_date": "Present", "four_tasks": ["Developed software applications using Python and JS"]}\n'
            '    ],\n'
            '    "education": [\n'
            '      {"degree": "B.Tech in Computer Science", "institution": "University Name", "description": "CGPA: 8.0/10", "start_date": "2024", "end_date": "2028"}\n'
            '    ],\n'
            '    "skills": {\n'
            '      "hard_skills": ["Python", "HTML", "CSS", "JavaScript"],\n'
            '      "soft_skills": ["Problem Solving", "Adaptability"]\n'
            '    }\n'
            '  }\n'
            "}\n"
        )

        messages = [SystemMessage(content=system_instruction)]

        # Append previous chat history
        for msg in chat_history[-10:]:  # Keep last 10 messages for context efficiency
            role = msg.get("role", "user")
            content = msg.get("content", "")
            if role == "user":
                messages.append(HumanMessage(content=content))
            elif role == "assistant":
                messages.append(AIMessage(content=content))

        # Current user turn
        current_content = user_message
        if extracted_doc_text:
            current_content += f"\n\n[DOCUMENT ATTACHMENT EXTRACTED TEXT]:\n{extracted_doc_text}"

        messages.append(HumanMessage(content=current_content))

        # Generate LLM response
        ai_response_msg = self.llm.invoke(messages)
        raw_text = ai_response_msg.content

        # Check if JSON block exists in response
        is_complete = False
        parsed_resume_data = None
        user_facing_response = raw_text

        if "```json" in raw_text:
            try:
                json_part = raw_text.split("```json")[1].split("```")[0].strip()
                raw_json = json.loads(json_part)

                # Robust Schema Normalization for UserInformation
                if "user_information" in raw_json:
                    user_info = raw_json["user_information"]
                    if not isinstance(user_info, dict):
                        user_info = {}
                        raw_json["user_information"] = user_info

                    # Map aliases
                    if "target_role" in user_info and "main_job_title" not in user_info:
                        user_info["main_job_title"] = user_info.pop("target_role")
                    if not user_info.get("main_job_title"):
                        user_info["main_job_title"] = "Software Engineer Intern"

                    if "summary" in user_info and "profile_description" not in user_info:
                        user_info["profile_description"] = user_info.pop("summary")
                    if not user_info.get("profile_description"):
                        user_info["profile_description"] = f"Motivated candidate seeking a position as {user_info.get('main_job_title', 'Software Engineer')}."

                    if "phone_number" in user_info and "phone" not in user_info:
                        user_info["phone"] = user_info.pop("phone_number")
                    if not user_info.get("phone"):
                        user_info["phone"] = "+1234567890"

                    if not user_info.get("email"):
                        user_info["email"] = "student@example.com"

                    if not user_info.get("name"):
                        user_info["name"] = "Candidate Name"

                    # Normalize experiences
                    exp_list = user_info.get("experiences")
                    if not isinstance(exp_list, list):
                        exp_list = []
                    
                    norm_exp = []
                    for e in exp_list:
                        if isinstance(e, dict):
                            jt = e.get("job_title") or e.get("title") or "Trainee"
                            comp = e.get("company") or "Academic Projects"
                            sd = e.get("start_date") or "2023"
                            ed = e.get("end_date") or "Present"
                            tasks = e.get("four_tasks")
                            if not isinstance(tasks, list):
                                desc = e.get("description") or "Contributed to software tasks and technical projects."
                                tasks = [desc]
                            norm_exp.append({
                                "job_title": jt,
                                "company": comp,
                                "location": e.get("location", "Remote"),
                                "start_date": sd,
                                "end_date": ed,
                                "four_tasks": tasks[:4]
                            })
                    user_info["experiences"] = norm_exp

                    # Normalize education
                    edu_list = user_info.get("education")
                    if not isinstance(edu_list, list):
                        edu_list = []

                    norm_edu = []
                    for ed in edu_list:
                        if isinstance(ed, dict):
                            inst = ed.get("institution") or "University"
                            deg = ed.get("degree") or "B.Tech"
                            s_date = ed.get("start_date") or "2024"
                            e_date = ed.get("end_date") or ed.get("year") or "2028"
                            desc = ed.get("description")
                            if not desc and ed.get("gpa"):
                                desc = f"GPA: {ed.get('gpa')}"
                            norm_edu.append({
                                "institution": inst,
                                "degree": deg,
                                "description": desc,
                                "start_date": s_date,
                                "end_date": e_date
                            })
                    user_info["education"] = norm_edu

                    # Normalize skills
                    skills_dict = user_info.get("skills")
                    if not isinstance(skills_dict, dict):
                        skills_dict = {}
                    hard_s = skills_dict.get("hard_skills") if isinstance(skills_dict.get("hard_skills"), list) else ["Python", "HTML", "CSS", "JavaScript"]
                    soft_s = skills_dict.get("soft_skills") if isinstance(skills_dict.get("soft_skills"), list) else ["Problem Solving", "Adaptability"]
                    user_info["skills"] = {"hard_skills": hard_s, "soft_skills": soft_s}

                    # Validate with Pydantic model
                    validated = ResumeData.parse_obj(raw_json)
                    parsed_resume_data = validated.model_dump()
                    is_complete = True

                    # User facing message
                    user_facing_response = raw_text.split("```json")[0].strip()
                    if not user_facing_response:
                        user_facing_response = "🎉 Excellent news! I've synthesized your details and generated your official ATS-optimized resume. Your downloadable PDF is ready below!"
            except Exception as e:
                logger.error(f"Error parsing generated resume JSON: {e}")

        return {
            "response": user_facing_response,
            "is_complete": is_complete,
            "resume_data": parsed_resume_data,
        }
