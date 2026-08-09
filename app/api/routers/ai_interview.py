"""AI Interview API router - booking and scheduling."""

import io
from typing import List, Dict, Optional
from datetime import datetime

from fastapi import APIRouter, HTTPException, Depends, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from app.database.repositories.user_repository import UserRepository
from app.database.repositories.settings_repository import SettingsRepository
from app.core.security import get_current_user, require_admin

ai_interview_router = APIRouter(prefix="/api/interview", tags=["interview"])


# ===== Pydantic Request Models =====
class BookInterviewRequest(BaseModel):
    resume_id: str
    phone_number: str
    interview_date: str
    interview_time: str
    target_company: Optional[str] = "Google"
    target_role: Optional[str] = "Software Engineer"


class ToggleServiceRequest(BaseModel):
    status: str  # 'active' or 'suspended'


# ===== Endpoints =====

@ai_interview_router.get("/status")
async def get_service_status():
    """Get global AI Interview service status."""
    settings_repo = SettingsRepository()
    srv_status = await settings_repo.get_setting("ai_interview_service", "active")
    return {"status": srv_status}


@ai_interview_router.post("/book")
async def book_interview(req: BookInterviewRequest, user_id: str = Depends(get_current_user)):
    """Schedule/book a new AI Interview."""
    settings_repo = SettingsRepository()
    srv_status = await settings_repo.get_setting("ai_interview_service", "active")
    if srv_status != "active":
        raise HTTPException(
            status_code=400, 
            detail="AI Interview Service is currently suspended by admin. Please try again later."
        )

    user_repo = UserRepository()
    user = await user_repo.get_user_by_id(user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    # Save to public.ai_interviews
    db_client = user_repo._get_supabase_client()
    try:
        payload = {
            "user_id": user_id,
            "user_name": user.get("name"),
            "phone_number": req.phone_number,
            "resume_id": req.resume_id,
            "interview_date": req.interview_date,
            "interview_time": req.interview_time,
            "target_company": req.target_company or "Company Mock",
            "target_role": req.target_role or "Software Engineer",
            "status": "pending",
            "created_at": datetime.now().isoformat()
        }
        res = db_client.table("ai_interviews").insert(payload).execute()
        if res.data:
            return {"success": True, "interview": res.data[0]}
        raise HTTPException(status_code=500, detail="Failed to save booking to database")
    except Exception as e:
        print(f"Error booking interview: {e}")
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")


@ai_interview_router.get("/mine")
async def get_my_interviews(user_id: str = Depends(get_current_user)):
    """Get all interviews booked by the logged-in user."""
    user_repo = UserRepository()
    db_client = user_repo._get_supabase_client()
    try:
        res = db_client.table("ai_interviews").select("*").eq("user_id", user_id).order("created_at", desc=True).execute()
        return res.data or []
    except Exception as e:
        print(f"Error fetching personal interviews: {e}")
        return []


@ai_interview_router.get("/admin/all")
async def get_all_interviews(user_id: str = Depends(get_current_user)):
    """Get all scheduled interviews across the system (admin only)."""
    user_repo = UserRepository()
    admin = await user_repo.get_user_by_id(user_id)
    if not admin or (not admin.get("is_admin", False) and admin.get("role") != "admin"):
        raise HTTPException(status_code=403, detail="Admin access required")

    db_client = user_repo._get_supabase_client()
    try:
        res = db_client.table("ai_interviews").select("*").order("created_at", desc=True).execute()
        return res.data or []
    except Exception as e:
        print(f"Error fetching all interviews: {e}")
        return []


@ai_interview_router.post("/admin/toggle")
async def toggle_service_status(req: ToggleServiceRequest, user_id: str = Depends(get_current_user)):
    """Toggle AI Interview service status between active and suspended (admin only)."""
    user_repo = UserRepository()
    admin = await user_repo.get_user_by_id(user_id)
    if not admin or (not admin.get("is_admin", False) and admin.get("role") != "admin"):
        raise HTTPException(status_code=403, detail="Admin access required")

    settings_repo = SettingsRepository()
    success = await settings_repo.set_setting("ai_interview_service", req.status)
    if success:
        return {"success": True, "status": req.status}
    raise HTTPException(status_code=500, detail="Failed to update settings in database")


@ai_interview_router.post("/{interview_id}/trigger")
async def trigger_interview(interview_id: str, user_id: str = Depends(get_current_user)):
    """Simulate triggering/running the AI Interview."""
    user_repo = UserRepository()
    user = await user_repo.get_user_by_id(user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    db_client = user_repo._get_supabase_client()
    try:
        # Check if interview exists
        res = db_client.table("ai_interviews").select("*").eq("id", interview_id).execute()
        if not res.data:
            raise HTTPException(status_code=404, detail="Interview booking not found")
        
        booking = res.data[0]
        # Only allow admin or the user who scheduled it
        if booking.get("user_id") != user_id and not user.get("is_admin", False) and user.get("role") != "admin":
            raise HTTPException(status_code=403, detail="Permission denied")

        # Update status to completed/triggered
        update_res = db_client.table("ai_interviews").update({"status": "completed"}).eq("id", interview_id).execute()
        if update_res.data:
            return {
                "success": True, 
                "message": "AI Interview session generated and mock calls placed successfully!",
                "booking": update_res.data[0]
            }
        raise HTTPException(status_code=500, detail="Failed to update booking status")
    except Exception as e:
        print(f"Error triggering interview: {e}")
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")


@ai_interview_router.delete("/admin/{interview_id}")
async def delete_interview(interview_id: str, user_id: str = Depends(get_current_user)):
    """Delete a scheduled interview booking (admin only)."""
    user_repo = UserRepository()
    admin = await user_repo.get_user_by_id(user_id)
    if not admin or (not admin.get("is_admin", False) and admin.get("role") != "admin"):
        raise HTTPException(status_code=403, detail="Admin access required")

    db_client = user_repo._get_supabase_client()
    try:
        res = db_client.table("ai_interviews").delete().eq("id", interview_id).execute()
        return {"success": True}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")


@ai_interview_router.get("/admin/export")
async def export_interviews_excel(user_id: str = Depends(get_current_user)):
    """Export all AI Interview bookings as an Excel (.xlsx) file (admin only)."""
    user_repo = UserRepository()
    admin = await user_repo.get_user_by_id(user_id)
    if not admin or (not admin.get("is_admin", False) and admin.get("role") != "admin"):
        raise HTTPException(status_code=403, detail="Admin access required")

    db_client = user_repo._get_supabase_client()
    try:
        res = db_client.table("ai_interviews").select("*").order("created_at", desc=True).execute()
        rows = res.data or []
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {str(e)}")

    # Build Excel using openpyxl
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed on server")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AI Interview Bookings"

    # Header row styling
    header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center")
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    headers = ["#", "Student Name", "Phone Number", "Interview Date", "Interview Time", "Status", "Registered At"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    # Data rows
    alt_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    for i, item in enumerate(rows, 1):
        created = item.get("created_at", "")
        if created:
            try:
                created = datetime.fromisoformat(created.replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M")
            except:
                pass
        row_data = [
            i,
            item.get("user_name") or "Unknown",
            item.get("phone_number") or "N/A",
            item.get("interview_date") or "N/A",
            item.get("interview_time") or "N/A",
            (item.get("status") or "pending").upper(),
            created
        ]
        ws.append(row_data)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=i + 1, column=col_idx)
            cell.border = thin
            cell.alignment = center
            if i % 2 == 0:
                cell.fill = alt_fill

    # Column widths
    col_widths = [5, 22, 18, 16, 14, 12, 20]
    for idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    # Stream as response
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"ai_interviews_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
